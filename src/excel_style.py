import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AUTHOR = "Krish Singh"
NAVY = "2F5496"

_THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11, name="Calibri")
BODY_FONT = Font(size=11, name="Calibri")


def set_author(workbook):
    workbook.properties.creator = AUTHOR
    workbook.properties.lastModifiedBy = AUTHOR


def style_sheet(ws, max_width=40, sample_rows=300):
    n_cols, n_rows = ws.max_column, ws.max_row

    for col in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    ws.row_dimensions[1].height = 30

    widths = [0] * (n_cols + 1)
    for row in ws.iter_rows(min_row=1, max_row=min(n_rows, sample_rows)):
        for cell in row:
            if cell.value is not None:
                widths[cell.column] = max(widths[cell.column], len(str(cell.value)))
    for col in range(1, n_cols + 1):
        ws.column_dimensions[get_column_letter(col)].width = max(10, min(max_width, widths[col] + 2))

    for row in ws.iter_rows(min_row=2, max_row=n_rows, max_col=n_cols):
        for cell in row:
            cell.font, cell.border = BODY_FONT, BORDER
            cell.alignment = Alignment(
                horizontal="right" if isinstance(cell.value, (int, float)) else "left", vertical="center",
            )
    ws.freeze_panes = "A2"


def sheet_name(name):
    return name[:31]


def save_workbook(path, sheets):
    """Write {sheet name: DataFrame} to an xlsx file with the standard header styling."""
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=sheet_name(name), index=False)

    workbook = openpyxl.load_workbook(path)
    for ws in workbook.worksheets:
        style_sheet(ws)
    set_author(workbook)
    workbook.save(path)
