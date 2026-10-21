import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from src.excel_style import AUTHOR

HEADER_FILL = "2F5496"
_BOLD = re.compile(r"\*\*(.+?)\*\*")


def _set_font(style, font):
    style.font.name = font
    fonts = style.element.rPr.rFonts
    fonts.set(qn("w:eastAsia"), font)
    for theme_attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(theme_attr), None)


def new_document(title, font="Calibri", size=11, heading_color=None):
    doc = Document()
    style = doc.styles["Normal"]
    _set_font(style, font)
    style.font.size = Pt(size)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        heading = doc.styles[name]
        _set_font(heading, font)
        if heading_color:
            heading.font.color.rgb = RGBColor.from_string(heading_color)
    doc.core_properties.author = AUTHOR
    doc.core_properties.last_modified_by = AUTHOR
    doc.core_properties.title = title
    return doc


def add_runs(paragraph, text, size=None, italic=False, color=None):
    parts = _BOLD.split(text)
    for i, part in enumerate(parts):
        if not part:
            continue
        run = paragraph.add_run(part)
        run.bold = i % 2 == 1
        run.italic = italic
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
    return paragraph


def para(doc, text, align=None, size=None, italic=False, color=None, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return add_runs(p, text, size=size, italic=italic, color=color)


def bullet(doc, text):
    return add_runs(doc.add_paragraph(style="List Bullet"), text)


def caption(doc, text):
    return para(doc, text, align="center", size=9, italic=True, space_after=4)


def _shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    props.append(shading)


def table(doc, headers, rows, size=9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(t.rows[0].cells, headers):
        cell.text = ""
        run = cell.paragraphs[0].add_run(str(text))
        run.bold = True
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor.from_string("FFFFFF")
        _shade(cell, HEADER_FILL)
    for row in rows:
        cells = t.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = ""
            run = cell.paragraphs[0].add_run("" if text is None else str(text))
            run.font.size = Pt(size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def pct(v, digits=1):
    return f"{float(v):.{digits}f}%"


def num(v):
    return f"{int(v):,}"


def p_value(p):
    return "< 0.001" if p < 0.001 else f"{p:.3f}"


def signed(v, digits=3):
    return f"{float(v):.{digits}f}".replace("-", "−")
