import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from src.config import COL_A1_NAME, OUTPUT_DIR
from src.excel_style import BORDER, HEADER_FILL, HEADER_FONT, set_author

OVERALL_FILE = OUTPUT_DIR / "Annotation_Worksheet.xlsx"
ABSA_FILE = OUTPUT_DIR / "ABSA_Annotation_Worksheet.xlsx"
SECOND_RATER_FILE = OUTPUT_DIR / "Annotation_Worksheet_rater2.xlsx"
TRAINING_FILE = OUTPUT_DIR / "Training_Worksheet.xlsx"
SHEET = "To Annotate"

ASPECT_COLUMNS = ["Timeliness", "Amount_of_Support", "Enrollment_Ease", "Official_Behavior", "Clarity_of_Info", "Overall_Impact"]
OVERALL_LABELS = ["Label", "Annotator_Initials", "Notes_optional"]
LEGACY_LABEL = "Human_Label"  # column name used before 2026-10-21
FILL_IN = PatternFill("solid", fgColor="FFF3C4")
LOCKED = PatternFill("solid", fgColor="F2F2F2")

LABEL_GUIDE = [
    "HOW TO LABEL (Label column):  shortcut - type 1 = Positive, 2 = Neutral, 3 = Negative",
    "  Positive  = mainly expresses satisfaction, benefit, relief or approval",
    "  Negative  = mainly expresses a problem, delay, frustration or complaint",
    "  Neutral   = factual with no clear tone, or genuinely mixed with no dominant tone",
    "",
    "Feedback_Text joins the farmer's three answers (positive aspects, problems, suggestions) with full stops.",
    "Answers that were only 'No' / 'Nothing' / 'Nil' are left out, so a short text can still be a complete response.",
]


def stratified_sample(pool, size, seed):
    fraction = size / len(pool)
    parts = [
        group.sample(n=min(len(group), max(1, round(len(group) * fraction))), random_state=seed)
        for _, group in pool.groupby("scheme_normalized")
    ]
    return pd.concat(parts).sample(frac=1, random_state=seed)


def worksheet_frame(sample, prefix, label_columns):
    frame = pd.DataFrame({
        "RespondentID": [f"{prefix}{str(i + 1).zfill(4)}" for i in range(len(sample))],
        "Source_Excel_Row": sample.index + 2,
        "Participant_Name": sample[COL_A1_NAME].values,
        "Scheme": sample["scheme_normalized"].values,
        "Feedback_Text": sample["scheme_feedback_text"].values,
    })
    for col in label_columns:
        frame[col] = None
    return frame


def has_labels(path, label_columns):
    if not path.exists():
        return False
    existing = pd.read_excel(path, sheet_name=SHEET)
    present = [c for c in [*label_columns, LEGACY_LABEL] if c in existing.columns]
    return bool(existing[present].notna().any().any())


def write_worksheet(path, frame, fill_in_cols, wide_cols, guide):
    workbook = openpyxl.Workbook()
    ws = workbook.active
    ws.title = SHEET
    ws.append(list(frame.columns))
    for row in frame.itertuples(index=False):
        ws.append([None if pd.isna(v) else v for v in row])

    for cell in ws[1]:
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    for idx, name in enumerate(frame.columns, start=1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = 62 if name in wide_cols else 15 if name in fill_in_cols else max(12, len(name) + 3)
        for row in range(2, len(frame) + 2):
            cell = ws[f"{letter}{row}"]
            cell.border = BORDER
            cell.fill = FILL_IN if name in fill_in_cols else LOCKED
            cell.alignment = Alignment(wrap_text=name in wide_cols, vertical="top")
    for row in range(2, len(frame) + 2):
        ws.row_dimensions[row].height = 46

    validation = DataValidation(type="list", formula1='"1,2,3,Positive,Neutral,Negative"', allow_blank=True)
    validation.error, validation.errorTitle = "Use 1 / 2 / 3 or Positive / Neutral / Negative", "Invalid label"
    ws.add_data_validation(validation)
    for name in fill_in_cols:
        letter = get_column_letter(list(frame.columns).index(name) + 1)
        validation.add(f"{letter}2:{letter}{len(frame) + 1}")

    readme = workbook.create_sheet("READ ME FIRST")
    for line in guide:
        readme.append([line])
    readme.column_dimensions["A"].width = 120
    readme["A1"].font = Font(bold=True)

    set_author(workbook)
    workbook.save(path)
